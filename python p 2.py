import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

# ==========================================
# EXCEL SETTINGS
# ==========================================

FILE_NAME = "hotel_data.xlsx"

HEADERS = [
    "Booking ID",
    "Guest Name",
    "Phone",
    "Room Number",
    "Room Type",
    "Check-in Date",
    "Check-out Date",
    "Guests",
    "Amount"
]


def create_excel():
    if not os.path.exists(FILE_NAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Bookings"
        sheet.append(HEADERS)
        workbook.save(FILE_NAME)
        workbook.close()


# ==========================================
# LOGIN
# ==========================================

def login():
    username = username_entry.get()
    password = password_entry.get()

    if username == "admin" and password == "1234":

        messagebox.showinfo(
            "Login",
            "Login Successful!"
        )

        login_window.destroy()
        open_dashboard()

    else:
        messagebox.showerror(
            "Login Failed",
            "Invalid Username or Password"
        )


# ==========================================
# DASHBOARD
# ==========================================

def open_dashboard():

    global dashboard

    dashboard = tk.Tk()
    dashboard.title("Hotel Management System")
    dashboard.geometry("1100x700")

    tk.Label(
        dashboard,
        text="HOTEL MANAGEMENT SYSTEM",
        font=("Arial", 24, "bold")
    ).pack(pady=30)

    button_frame = tk.Frame(dashboard)
    button_frame.pack()

    tk.Button(
        button_frame,
        text="Add Booking",
        width=18,
        command=open_add_booking
    ).grid(row=0, column=0, padx=10)

    tk.Button(
        button_frame,
        text="View Records",
        width=18,
        command=view_records
    ).grid(row=0, column=1, padx=10)

    tk.Button(
        button_frame,
        text="Logout",
        width=18,
        command=logout
    ).grid(row=0, column=2, padx=10)

    dashboard.mainloop()


# ==========================================
# ADD BOOKING
# ==========================================

def open_add_booking():

    add_window = tk.Toplevel(dashboard)

    add_window.title("Add Booking")
    add_window.geometry("520x620")

    tk.Label(
        add_window,
        text="ADD HOTEL BOOKING",
        font=("Arial", 20, "bold")
    ).pack(pady=20)

    form = tk.Frame(add_window)
    form.pack()

    entries = {}

    for i, label in enumerate(HEADERS):

        tk.Label(
            form,
            text=label
        ).grid(
            row=i,
            column=0,
            padx=10,
            pady=7,
            sticky="w"
        )

        entry = tk.Entry(
            form,
            width=30
        )

        entry.grid(
            row=i,
            column=1,
            padx=10,
            pady=7
        )

        entries[label] = entry

    def clear_fields():

        for entry in entries.values():
            entry.delete(0, tk.END)

    def save_booking():

        data = []

        for label in HEADERS:
            data.append(entries[label].get().strip())

        if data[0] == "":
            messagebox.showwarning(
                "Warning",
                "Booking ID is required."
            )
            return

        if data[1] == "":
            messagebox.showwarning(
                "Warning",
                "Guest Name is required."
            )
            return

        if data[3] == "":
            messagebox.showwarning(
                "Warning",
                "Room Number is required."
            )
            return

        create_excel()

        workbook = load_workbook(FILE_NAME)
        sheet = workbook.active

        # Check duplicate Booking ID
        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            if str(row[0]) == str(data[0]):

                messagebox.showerror(
                    "Error",
                    "Booking ID already exists."
                )

                workbook.close()
                return

        sheet.append(data)

        workbook.save(FILE_NAME)
        workbook.close()

        messagebox.showinfo(
            "Success",
            "Booking saved successfully!"
        )

        clear_fields()

    tk.Button(
        add_window,
        text="Save Booking",
        width=18,
        command=save_booking
    ).pack(pady=15)

    tk.Button(
        add_window,
        text="Clear / Reset",
        width=18,
        command=clear_fields
    ).pack()


# ==========================================
# VIEW RECORDS
# ==========================================

def view_records():

    view_window = tk.Toplevel(dashboard)

    view_window.title("View Hotel Records")
    view_window.geometry("1250x650")

    tk.Label(
        view_window,
        text="HOTEL BOOKING RECORDS",
        font=("Arial", 20, "bold")
    ).pack(pady=15)

    # Search section
    search_frame = tk.Frame(view_window)
    search_frame.pack(pady=5)

    tk.Label(
        search_frame,
        text="Search:"
    ).pack(side="left")

    search_entry = tk.Entry(
        search_frame,
        width=30
    )

    search_entry.pack(
        side="left",
        padx=10
    )

    # Table
    table_frame = tk.Frame(view_window)
    table_frame.pack(
        fill="both",
        expand=True,
        padx=10,
        pady=10
    )

    tree = ttk.Treeview(
        table_frame,
        columns=HEADERS,
        show="headings"
    )

    for column in HEADERS:

        tree.heading(
            column,
            text=column
        )

        tree.column(
            column,
            width=120
        )

    tree.pack(
        side="left",
        fill="both",
        expand=True
    )

    scrollbar = ttk.Scrollbar(
        table_frame,
        orient="vertical",
        command=tree.yview
    )

    scrollbar.pack(
        side="right",
        fill="y"
    )

    tree.configure(
        yscrollcommand=scrollbar.set
    )

    # ======================================
    # LOAD RECORDS
    # ======================================

    def load_records():

        for item in tree.get_children():
            tree.delete(item)

        create_excel()

        workbook = load_workbook(FILE_NAME)
        sheet = workbook.active

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            tree.insert(
                "",
                tk.END,
                values=row
            )

        workbook.close()

    # ======================================
    # SEARCH
    # ======================================

    def search_records():

        search_text = search_entry.get().lower()

        for item in tree.get_children():
            tree.delete(item)

        workbook = load_workbook(FILE_NAME)
        sheet = workbook.active

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            row_text = " ".join(
                str(value).lower()
                for value in row
                if value is not None
            )

            if search_text in row_text:

                tree.insert(
                    "",
                    tk.END,
                    values=row
                )

        workbook.close()

    # ======================================
    # UPDATE RECORD
    # ======================================

    def update_record():

        selected = tree.selection()

        if not selected:

            messagebox.showwarning(
                "Select Record",
                "Please select a record first."
            )

            return

        values = tree.item(
            selected[0],
            "values"
        )

        update_window = tk.Toplevel(view_window)

        update_window.title("Update Booking")
        update_window.geometry("520x620")

        tk.Label(
            update_window,
            text="UPDATE BOOKING",
            font=("Arial", 20, "bold")
        ).pack(pady=20)

        form = tk.Frame(update_window)
        form.pack()

        update_entries = {}

        for i, label in enumerate(HEADERS):

            tk.Label(
                form,
                text=label
            ).grid(
                row=i,
                column=0,
                padx=10,
                pady=7,
                sticky="w"
            )

            entry = tk.Entry(
                form,
                width=30
            )

            entry.grid(
                row=i,
                column=1,
                padx=10,
                pady=7
            )

            entry.insert(0, values[i])

            update_entries[label] = entry

        def save_update():

            new_data = []

            for label in HEADERS:

                new_data.append(
                    update_entries[label].get().strip()
                )

            if new_data[0] == "":
                messagebox.showwarning(
                    "Warning",
                    "Booking ID is required."
                )
                return

            workbook = load_workbook(FILE_NAME)
            sheet = workbook.active

            for row in range(
                2,
                sheet.max_row + 1
            ):

                if str(
                    sheet.cell(row, 1).value
                ) == str(values[0]):

                    for column in range(
                        1,
                        len(HEADERS) + 1
                    ):

                        sheet.cell(
                            row,
                            column
                        ).value = new_data[column - 1]

                    break

            workbook.save(FILE_NAME)
            workbook.close()

            messagebox.showinfo(
                "Success",
                "Booking updated successfully!"
            )

            update_window.destroy()

            load_records()

        tk.Button(
            update_window,
            text="Update Booking",
            width=18,
            command=save_update
        ).pack(pady=15)

    # ======================================
    # DELETE RECORD
    # ======================================

    def delete_record():

        selected = tree.selection()

        if not selected:

            messagebox.showwarning(
                "Select Record",
                "Please select a record first."
            )

            return

        values = tree.item(
            selected[0],
            "values"
        )

        booking_id = values[0]

        answer = messagebox.askyesno(
            "Confirm Delete",
            "Are you sure you want to delete Booking ID "
            + str(booking_id)
            + "?"
        )

        if answer:

            workbook = load_workbook(FILE_NAME)
            sheet = workbook.active

            for row in range(
                2,
                sheet.max_row + 1
            ):

                if str(
                    sheet.cell(row, 1).value
                ) == str(booking_id):

                    sheet.delete_rows(
                        row,
                        1
                    )

                    break

            workbook.save(FILE_NAME)
            workbook.close()

            messagebox.showinfo(
                "Deleted",
                "Booking deleted successfully!"
            )

            load_records()

    # ======================================
    # BUTTONS
    # ======================================

    tk.Button(
        search_frame,
        text="Search",
        width=12,
        command=search_records
    ).pack(side="left")

    tk.Button(
        search_frame,
        text="View All",
        width=12,
        command=load_records
    ).pack(
        side="left",
        padx=5
    )

    tk.Button(
        search_frame,
        text="Update",
        width=12,
        command=update_record
    ).pack(
        side="left",
        padx=5
    )

    tk.Button(
        search_frame,
        text="Delete",
        width=12,
        command=delete_record
    ).pack(
        side="left",
        padx=5
    )

    load_records()


# ==========================================
# LOGOUT
# ==========================================

def logout():

    dashboard.destroy()

    open_login()


# ==========================================
# LOGIN WINDOW
# ==========================================

def open_login():

    global login_window
    global username_entry
    global password_entry

    login_window = tk.Tk()

    login_window.title(
        "Hotel Management System - Login"
    )

    login_window.geometry("450x350")

    tk.Label(
        login_window,
        text="HOTEL MANAGEMENT SYSTEM",
        font=("Arial", 18, "bold")
    ).pack(pady=30)

    tk.Label(
        login_window,
        text="Username"
    ).pack()

    username_entry = tk.Entry(
        login_window,
        width=30
    )

    username_entry.pack(pady=5)

    tk.Label(
        login_window,
        text="Password"
    ).pack()

    password_entry = tk.Entry(
        login_window,
        width=30,
        show="*"
    )

    password_entry.pack(pady=5)

    tk.Button(
        login_window,
        text="Login",
        width=15,
        command=login
    ).pack(pady=20)

    tk.Label(
        login_window,
        text="Demo Login: admin / 1234"
    ).pack()

    login_window.mainloop()


# ==========================================
# START PROJECT
# ==========================================

create_excel()

open_login()
